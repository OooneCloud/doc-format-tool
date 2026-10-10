# -*- coding: utf-8 -*-
"""文档格式转换工具 —— 免费、离线、本地运行

功能：
  1. Markdown(.md) → Word(.docx)
  2. Markdown(.md) → PDF
  3. PDF → 图片（PNG / JPG，可调 DPI）
  支持 MD 中的本地图片与网络图片，转换时自动嵌入。

依赖（首次使用前安装一次即可）：
  pip install pypandoc-binary xhtml2pdf pymupdf markdown python-docx sv-ttk pillow
"""

import os
import re
import sys
import queue
import hashlib
import logging
import shutil
import tempfile
import threading
import traceback
import urllib.request
from urllib.parse import urlparse
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP_TITLE = "文档格式转换工具"
# 打包成 exe 后，把 error.log / reference.docx 等放在 exe 旁边而不是临时目录
if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))
ERROR_LOG = os.path.join(APP_DIR, "error.log")

WINDIR = os.environ.get("WINDIR", r"C:\Windows")
FONT_DIR = os.path.join(WINDIR, "Fonts")
FONT_REGULAR_CANDIDATES = ["Deng.ttf", "simfang.ttf", "simhei.ttf"]
FONT_BOLD_CANDIDATES = ["Dengb.ttf", "simhei.ttf"]

# ---------------- 界面配色（浅色 Fluent 风格） ----------------

FONT = "Microsoft YaHei UI"
C_WINDOW = "#f3f3f3"        # 窗口底色
C_SIDEBAR = "#fafafa"       # 侧栏底色
C_CARD = "#ffffff"          # 卡片底色
C_BORDER = "#e4e4e4"        # 边框
C_HOVER = "#ececec"         # 悬停
C_ACCENT = "#0f6cbd"        # 主色（蓝）
C_ACCENT_HOVER = "#115ea3"
C_ACCENT_LIGHT = "#e8f1fb"  # 浅蓝（选中项 / 次要按钮）
C_ACCENT_SOFT_HOVER = "#d8e7f7"
C_TEXT = "#1b1b1b"
C_SUB = "#6f6f6f"           # 次要文字
C_OK = "#0f7b0f"
C_ERR = "#c50f1f"
C_DISABLED = "#b8b8b8"

TARGET_NAME = {"docx": "Word", "pdf": "PDF", "img": "图片"}
IMG_MAX_W = 620  # PDF 页面内容区最大图片宽度（px）
UI_SHRINK = 0.75  # 整体界面缩放（按原比例缩小 1/4）

# 拖放支持（可选：未安装 windnd 时功能正常，只是不能拖文件）
try:
    import windnd
    _HAS_DND = True
except Exception:
    _HAS_DND = False


def _find_font_file(names):
    for n in names:
        p = os.path.join(FONT_DIR, n)
        if os.path.exists(p):
            return p
    return None


def read_text_auto(path):
    """读取文本文件，自动尝试 utf-8 / gb 系编码"""
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# ---------------- 转换核心（不依赖界面，便于独立测试） ----------------

REFERENCE_DOCX = os.path.join(APP_DIR, "reference_v2.docx")  # v2: Markdown 编辑器风格
DOCX_FONT_LATIN = "DengXian"
DOCX_FONT_EA = "等线"
DOCX_FONT_MONO = "Consolas"  # 代码块/行内代码的等宽字体

# OOXML 子元素规范顺序（保证插入位置合法，避免 Word 报文档损坏）
_PPR_ORDER = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
              "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
              "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
              "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
              "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
              "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
              "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_RPR_ORDER = ("w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
              "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow",
              "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
              "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position",
              "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
              "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
              "w:eastAsianLayout", "w:specVanish", "w:oMath")
_TCPR_ORDER = ("w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge",
               "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
               "w:tcFitText", "w:vAlign", "w:hideMark")
_TBLPR_ORDER = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
                "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
                "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
                "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


def _oxml_insert(parent, el, order):
    """按 OOXML 规定的子元素顺序插入"""
    key = "w:" + el.tag.rsplit("}", 1)[-1]
    if key not in order:
        parent.append(el)
        return
    idx = order.index(key)
    for child in parent:
        ctag = "w:" + child.tag.rsplit("}", 1)[-1]
        if ctag in order and order.index(ctag) > idx:
            child.addprevious(el)
            return
    parent.append(el)


def _oxml_new(tag, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), str(v))
    return el


def md_file_to_docx(md_path, docx_path, log=print):
    """单个 md → docx，由 pandoc 完成转换，支持表格/代码块/图片/嵌套列表。
    本地图片（相对/绝对路径）与网络图片都会嵌入文档。
    挂载定制的 reference.docx，让 Word 排版与 PDF 风格一致。"""
    import pypandoc
    src_dir = os.path.dirname(os.path.abspath(md_path))
    extra = ["--resource-path=%s" % src_dir]
    ref = _ensure_reference_docx(log=log)
    if ref:
        extra.append("--reference-doc=%s" % ref)
    pypandoc.convert_file(
        md_path, "docx",
        format="markdown+east_asian_line_breaks",
        outputfile=docx_path,
        extra_args=extra,
    )
    try:
        _style_docx_tables(docx_path)
    except Exception:
        log("    提示：表格样式微调失败（不影响内容），详见 error.log")
        try:
            with open(ERROR_LOG, "a", encoding="utf-8") as ef:
                ef.write("[表格样式] %s\n%s\n" % (docx_path, traceback.format_exc()))
        except Exception:
            pass


def _ensure_reference_docx(log=print):
    """生成（或复用）pandoc 参考样式文档：等线正文、深色分级标题、
    灰底代码块、引用块竖线、A4 页面。失败时返回 None（退回 pandoc 默认样式）。
    程序更新后（代码比缓存新）会自动重建模板。"""
    if os.path.exists(REFERENCE_DOCX):
        try:
            code_path = sys.executable if getattr(sys, "frozen", False) \
                else os.path.abspath(__file__)
            if os.path.getmtime(REFERENCE_DOCX) >= os.path.getmtime(code_path):
                return REFERENCE_DOCX
        except OSError:
            return REFERENCE_DOCX
    try:
        import io
        import pypandoc
        import docx
        from docx.oxml.ns import qn
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.style import WD_STYLE_TYPE

        pandoc = pypandoc.get_pandoc_path()
        import subprocess
        proc = subprocess.run([pandoc, "--print-default-data-file", "reference.docx"],
                              capture_output=True, timeout=60)
        if proc.returncode != 0 or not proc.stdout.startswith(b"PK"):
            return None

        def set_font(st, size=None, color=None, bold=None, latin=None):
            latin = latin or DOCX_FONT_LATIN
            f = st.font
            if size is not None:
                f.size = Pt(size)
            if color is not None:
                f.color.rgb = RGBColor.from_string(color)
            if bold is not None:
                f.bold = bold
            f.name = latin
            rpr = st.element.get_or_add_rPr()
            rf = rpr.get_or_add_rFonts()
            # 主题字体优先级高于显式字体，必须移除
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(a)) is not None:
                    del rf.attrib[qn(a)]
            rf.set(qn("w:ascii"), latin)
            rf.set(qn("w:hAnsi"), latin)
            rf.set(qn("w:eastAsia"), DOCX_FONT_EA)

        def get_cstyle(doc, name):
            """取/建字符样式（语法高亮 token 用）"""
            try:
                return doc.styles[name]
            except KeyError:
                st = doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
                st.quick_style = True
                return st

        def set_char_color(doc, name, color):
            st = get_cstyle(doc, name)
            st.font.color.rgb = RGBColor.from_string(color)

        def get_pstyle(doc, name):
            """取样式，不存在则创建（pandoc 按名称匹配，新建同样生效）"""
            try:
                return doc.styles[name]
            except KeyError:
                st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                st.quick_style = True
                try:
                    st.base_style = doc.styles["Normal"]
                except KeyError:
                    pass
                return st

        bio = io.BytesIO(proc.stdout)
        doc = docx.Document(bio)

        # 页面：A4，边距与 PDF 一致（上下 2.2cm，左右 2cm）
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2)

        # 正文（颜色取 Markdown 编辑器的正文色）
        set_font(doc.styles["Normal"], size=11, color="24292E")
        for name in ("Body Text", "First Paragraph", "Compact"):
            st = get_pstyle(doc, name)
            set_font(st, size=11, color="24292E")
            if name != "Compact":
                st.paragraph_format.line_spacing = 1.25
                st.paragraph_format.space_after = Pt(8)

        # 标题：深色加粗，1/2 级加浅灰下划线（GitHub README 风格）
        sizes = {"Heading 1": 20, "Heading 2": 16, "Heading 3": 13,
                 "Heading 4": 12, "Heading 5": 12, "Heading 6": 12}
        for name, size in sizes.items():
            st = get_pstyle(doc, name)
            set_font(st, size=size, color="24292E", bold=True)
            st.paragraph_format.space_before = Pt(16)
            st.paragraph_format.space_after = Pt(8)
            st.paragraph_format.line_spacing = 1.3
            if name in ("Heading 1", "Heading 2"):
                ppr = st.element.get_or_add_pPr()
                pbdr = _oxml_new("w:pBdr")
                pbdr.append(_oxml_new("w:bottom", **{"w:val": "single", "w:sz": 6,
                                                     "w:space": 4, "w:color": "D0D7DE"}))
                _oxml_insert(ppr, pbdr, _PPR_ORDER)

        # 代码块：等宽字体 + 浅灰底 + 细边框（Markdown 编辑器风格）
        st = get_pstyle(doc, "Source Code")
        set_font(st, size=10, latin=DOCX_FONT_MONO)
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.space_before = Pt(6)
        st.paragraph_format.space_after = Pt(6)
        ppr = st.element.get_or_add_pPr()
        _oxml_insert(ppr, _oxml_new("w:shd", **{"w:val": "clear",
                                                "w:color": "auto",
                                                "w:fill": "F6F8FA"}), _PPR_ORDER)
        pbdr = _oxml_new("w:pBdr")
        for edge in ("top", "left", "bottom", "right"):
            pbdr.append(_oxml_new("w:" + edge, **{"w:val": "single", "w:sz": 4,
                                                  "w:space": 4, "w:color": "D0D7DE"}))
        _oxml_insert(ppr, pbdr, _PPR_ORDER)

        # 语法高亮配色（GitHub light 主题）
        for names, color in (
                (("KeywordTok", "DataTypeTok", "ControlFlowTok", "ImportTok"), "D73A49"),
                (("StringTok", "CharTok", "VerbatimStringTok",
                  "SpecialStringTok", "SpecialCharTok"), "032F62"),
                (("CommentTok", "DocumentationTok", "CommentVarTok"), "6A737D"),
                (("FunctionTok", "PreProcTok"), "6F42C1"),
                (("DecValTok", "FloatTok", "BaseNTok", "ConstantTok",
                  "BuiltInTok", "AttributeTok"), "005CC5"),
                (("VariableTok",), "E36209"),
                (("OperatorTok", "NormalTok", "OtherTok", "InformationTok",
                  "ExtensionTok", "AlertTok", "WarningTok", "ErrorTok"), "24292E")):
            for n in names:
                set_char_color(doc, n, color)

        # 行内代码：等宽字体 + 浅灰底纹 + 深红色（Markdown 风格）
        try:
            st = doc.styles["Verbatim Char"]
            set_font(st, size=10, color="C7254E", latin=DOCX_FONT_MONO)
            rpr = st.element.get_or_add_rPr()
            _oxml_insert(rpr, _oxml_new("w:shd", **{"w:val": "clear",
                                                    "w:color": "auto",
                                                    "w:fill": "EFF1F3"}), _RPR_ORDER)
        except KeyError:
            pass

        # 超链接：GitHub 标准蓝、无下划线
        try:
            st = doc.styles["Hyperlink"]
            st.font.color.rgb = RGBColor.from_string("0969DA")
            rpr = st.element.get_or_add_rPr()
            u = rpr.find(qn("w:u"))
            if u is not None:
                rpr.remove(u)
        except KeyError:
            pass

        # 引用块：左侧竖线 + 灰字 + 缩进
        st = get_pstyle(doc, "Block Text")
        set_font(st, size=11, color="57606A")
        st.paragraph_format.left_indent = Cm(0.5)
        ppr = st.element.get_or_add_pPr()
        pbdr = _oxml_new("w:pBdr")
        pbdr.append(_oxml_new("w:left", **{"w:val": "single", "w:sz": 24,
                                           "w:space": 4, "w:color": "D0D7DE"}))
        _oxml_insert(ppr, pbdr, _PPR_ORDER)

        # 兜底：所有段落/字符样式的中文字体统一为等线
        for st in doc.styles:
            try:
                if st.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
                    rpr = st.element.get_or_add_rPr()
                    rf = rpr.get_or_add_rFonts()
                    if rf.get(qn("w:eastAsia")) is None:
                        rf.set(qn("w:eastAsia"), DOCX_FONT_EA)
            except Exception:
                continue

        doc.save(REFERENCE_DOCX)
        log("    提示：首次运行，已生成 Word 样式模板（Markdown 风格）")
        return REFERENCE_DOCX
    except Exception:
        try:
            with open(ERROR_LOG, "a", encoding="utf-8") as f:
                f.write("[样式模板] \n%s\n" % traceback.format_exc())
        except Exception:
            pass
        return None


def _style_docx_tables(docx_path):
    """转换后微调：表格（全边框、表头灰底跨页重复）+ 分隔线（浅灰细线）"""
    from docx.oxml.ns import qn
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = docx.Document(docx_path)
    # 分隔线：pandoc 生成黑色 VML 粗线，改成 Markdown 风格的浅灰细线
    ns_v = "{urn:schemas-microsoft-com:vml}"
    for rect in list(doc.element.body.iter(ns_v + "rect")):
        p = rect
        while p is not None and p.tag != qn("w:p"):
            p = p.getparent()
        if p is None:
            continue
        for child in list(p):
            if child.tag != qn("w:pPr"):
                p.remove(child)
        ppr = p.get_or_add_pPr()
        old = ppr.find(qn("w:pBdr"))
        if old is not None:
            ppr.remove(old)
        pbdr = _oxml_new("w:pBdr")
        pbdr.append(_oxml_new("w:bottom", **{"w:val": "single", "w:sz": 6,
                                             "w:space": 1, "w:color": "D0D7DE"}))
        _oxml_insert(ppr, pbdr, _PPR_ORDER)
        _oxml_insert(ppr, _oxml_new("w:spacing", **{"w:before": 160, "w:after": 160,
                                                    "w:line": 240,
                                                    "w:lineRule": "auto"}), _PPR_ORDER)
    for table in doc.tables:
        tblPr = table._tbl.tblPr
        # 全边框（0.5pt 灰）
        old = tblPr.find(qn("w:tblBorders"))
        if old is not None:
            tblPr.remove(old)
        borders = _oxml_new("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(_oxml_new("w:" + edge, **{"w:val": "single", "w:sz": 4,
                                                     "w:space": 0, "w:color": "D0D7DE"}))
        _oxml_insert(tblPr, borders, _TBLPR_ORDER)
        # 单元格内边距（上下 3pt，左右 6pt）
        old = tblPr.find(qn("w:tblCellMar"))
        if old is not None:
            tblPr.remove(old)
        mar = _oxml_new("w:tblCellMar")
        for side, wd in (("top", 60), ("left", 120), ("bottom", 60), ("right", 120)):
            mar.append(_oxml_new("w:" + side, **{"w:w": wd, "w:type": "dxa"}))
        _oxml_insert(tblPr, mar, _TBLPR_ORDER)
        # 表头行
        if not table.rows:
            continue
        first = table.rows[0]
        trpr = first._tr.get_or_add_trPr()
        if trpr.find(qn("w:tblHeader")) is None:
            trpr.append(_oxml_new("w:tblHeader"))
        for cell in first.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            if tcpr.find(qn("w:shd")) is None:
                _oxml_insert(tcpr, _oxml_new("w:shd", **{"w:val": "clear",
                                                         "w:color": "auto",
                                                         "w:fill": "F6F8FA"}),
                             _TCPR_ORDER)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.bold = True
        # 列宽按内容长度自动分配（长内容的列更宽）
        import unicodedata

        def vis_len(s):
            return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
                       for ch in s)

        try:
            ncols = len(table.columns)
            grid = table._tbl.find(qn("w:tblGrid"))
            cols = grid.findall(qn("w:gridCol")) if grid is not None else []
            if ncols and len(cols) == ncols:
                content_w = 9639  # A4 宽 21cm - 左右边距各 2cm，单位 twips
                weights = []
                for c in range(ncols):
                    wmax = 0
                    for r in range(len(table.rows)):
                        try:
                            t = "".join(p.text for p in table.cell(r, c).paragraphs).strip()
                        except Exception:
                            t = ""
                        wmax = max(wmax, min(vis_len(t), 60))
                    weights.append(max(wmax, 6))  # 下限，避免窄列被压没
                total = sum(weights)
                widths = [max(int(content_w * x / total), 700) for x in weights]
                scale = content_w / sum(widths)
                widths = [int(x * scale) for x in widths]
                # 表格整体宽度 + 允许 Word 按内容再微调
                old = tblPr.find(qn("w:tblW"))
                if old is not None:
                    tblPr.remove(old)
                _oxml_insert(tblPr, _oxml_new("w:tblW", **{"w:w": content_w,
                                                           "w:type": "dxa"}), _TBLPR_ORDER)
                old = tblPr.find(qn("w:tblLayout"))
                if old is not None:
                    tblPr.remove(old)
                _oxml_insert(tblPr, _oxml_new("w:tblLayout", **{"w:type": "autofit"}),
                             _TBLPR_ORDER)
                for col, wd in zip(cols, widths):
                    col.set(qn("w:w"), str(wd))
                for row in table.rows:
                    for c_idx, cell in enumerate(row.cells):
                        if c_idx >= ncols:
                            continue
                        tcpr = cell._tc.get_or_add_tcPr()
                        old = tcpr.find(qn("w:tcW"))
                        if old is not None:
                            tcpr.remove(old)
                        _oxml_insert(tcpr, _oxml_new("w:tcW", **{"w:w": widths[c_idx],
                                                                 "w:type": "dxa"}),
                                     _TCPR_ORDER)
        except Exception:
            pass
    doc.save(docx_path)


def _prep_images(body, md_dir, log=print):
    """处理 HTML 中的 <img>：
    1) 网络图片下载到临时目录，替换为本地路径（失败则保留原样并提示）；
    2) 超过页面宽度的图片加 width 属性，避免溢出。
    返回 (新body, 临时目录)"""
    tmpdir = tempfile.mkdtemp(prefix="doctools_img_")
    errors = []

    def fetch(url):
        ext = os.path.splitext(urlparse(url).path)[1]
        if len(ext) > 6 or not ext:
            ext = ".png"
        path = os.path.join(tmpdir, hashlib.md5(url.encode("utf-8")).hexdigest() + ext)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = resp.read()
        if not data:
            raise RuntimeError("空文件")
        with open(path, "wb") as f:
            f.write(data)
        return path

    def fix(m):
        tag = m.group(0)
        sm = re.search(r'src=["\']([^"\']+)["\']', tag)
        if not sm:
            return tag
        src = sm.group(1)
        if re.match(r"^https?://", src, re.I):
            try:
                src = fetch(src)
            except Exception as e:
                errors.append("%s（%s）" % (src[:70], e))
                return tag
            tag = re.sub(r'src=["\'][^"\']+["\']',
                         'src="%s"' % src.replace("\\", "/"), tag, count=1)
        else:
            src = src if os.path.isabs(src) else os.path.normpath(os.path.join(md_dir, src))
        # 过宽的图片限制宽度（只测尺寸，不改源文件）
        try:
            from PIL import Image
            with Image.open(src) as im:
                w, _h = im.size
            if w > IMG_MAX_W and not re.search(r"\swidth=", tag, re.I):
                tag = re.sub(r"\s*/?>$", ' width="%d">' % IMG_MAX_W, tag)
        except Exception:
            pass
        return tag

    new_body = re.sub(r"<img[^>]*>", fix, body, flags=re.I)
    for e in errors:
        log("    提示：网络图片下载失败，已跳过：%s" % e)
    return new_body, tmpdir


def _md_to_html(md_path, log=print):
    import markdown as mdlib
    text = read_text_auto(md_path)
    # 去掉 YAML front matter（文件开头 --- 包裹的元数据）
    text = re.sub(r"\A---\s*\n.*?\n---\s*\n", "", text, flags=re.S)
    return mdlib.markdown(
        text,
        extensions=["tables", "fenced_code", "sane_lists"],
        output_format="xhtml",
    )


def _register_cjk_fonts():
    """用 reportlab 直接注册中文字体（比 xhtml2pdf 的 @font-face 可靠）。
    返回 (正文字体名, 粗体字体名)，找不到可用字体时返回 (None, None)。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    def try_register(name, candidates):
        if name in pdfmetrics.getRegisteredFontNames():
            return name
        for n in candidates:
            p = os.path.join(FONT_DIR, n)
            if os.path.exists(p):
                try:
                    pdfmetrics.registerFont(TTFont(name, p))
                    return name
                except Exception:
                    continue
        return None

    reg = try_register("CNBody", FONT_REGULAR_CANDIDATES)
    bold = try_register("CNBold", FONT_BOLD_CANDIDATES)
    return reg, bold


def _sanitize_links(body):
    """把 <a> 里非网络链接（文内锚点 #xxx、相对路径等）展开为普通文本：
    PDF 引擎无法为这类链接生成书签目标，会直接报 format not resolved。
    http/https/mailto 链接保留可点击。返回 (新body, 展开的链接数)。"""
    count = [0]

    def repl(m):
        tag, inner = m.group(1), m.group(2)
        hm = re.search(r'href=["\']([^"\']*)["\']', tag, re.I)
        href = hm.group(1) if hm else ""
        if href and re.match(r"^(https?|mailto):", href, re.I):
            return m.group(0)
        count[0] += 1
        return inner

    new = re.sub(r"<a\b([^>]*)>(.*?)</a>", repl, body, flags=re.I | re.S)
    return new, count[0]


def md_file_to_pdf(md_path, pdf_path, log=print):
    """单个 md → pdf：先转 HTML，再用 xhtml2pdf 排版为 A4 页面"""
    from xhtml2pdf import pisa
    from xhtml2pdf import default as pisa_default
    logging.getLogger("xhtml2pdf").setLevel(logging.ERROR)

    md_dir = os.path.dirname(os.path.abspath(md_path))
    body = _md_to_html(md_path, log=log)
    body, n_links = _sanitize_links(body)
    if n_links:
        log("    提示：%d 个文内跳转链接已转为普通文字" % n_links)
    body, tmpdir = _prep_images(body, md_dir, log=log)
    try:
        reg, bold = _register_cjk_fonts()

        # xhtml2pdf 只认自己 DEFAULT_FONT 字典里的字体名，
        # 需把 reportlab 注册好的字体注入进去（键必须全小写）
        body_family = "Helvetica"
        head_family = "Helvetica-Bold"
        if reg:
            pisa_default.DEFAULT_FONT[reg.lower()] = reg
            body_family = reg
            head_family = bold or reg
        if bold:
            pisa_default.DEFAULT_FONT[bold.lower()] = bold

        css = """
@page { size: a4 portrait; margin: 2.2cm 2cm; }
body, p, td, th, li, dt, dd, div, blockquote, pre, code,
h1, h2, h3, h4, h5, h6 { -pdf-word-wrap: CJK; }
body { font-family: %(body)s; font-size: 11pt; line-height: 170%%; color: #1f1f1f; }
h1, h2, h3, h4, h5, h6 { font-family: %(head)s; color: #111111; margin-top: 14pt; margin-bottom: 6pt; }
h1 { font-size: 20pt; }
h2 { font-size: 16pt; }
h3 { font-size: 13pt; }
h4, h5, h6 { font-size: 12pt; }
strong, b, th { font-family: %(head)s; }
pre, code, pre code { font-family: %(body)s; }
pre { font-size: 9.5pt; background-color: #f6f6f6;
      border: 0.5pt solid #cccccc; padding: 6pt; }
code { color: #b03030; }
blockquote { border-left: 3pt solid #bbbbbb; padding-left: 8pt; color: #555555; margin-left: 6pt; }
table { margin: 8pt 0; }
th, td { border: 0.5pt solid #999999; padding: 3pt 6pt; font-size: 10pt; }
th { background-color: #ececec; }
a { color: #1a5fb4; }
hr { color: #cccccc; }
img { margin: 6pt 0; }
""" % {"body": body_family, "head": head_family}

        html = ("<html><head><meta charset='utf-8'><style>%s</style></head>"
                "<body>%s</body></html>" % (css, body))

        def link_callback(uri, *args):
            """把相对路径图片解析到 md 所在目录（兼容中文/空格文件名）"""
            import urllib.parse
            if re.match(r"^[a-zA-Z]+:", uri) or uri.startswith("#"):
                return uri
            path = urllib.parse.unquote(uri)
            path = path if os.path.isabs(path) else os.path.join(md_dir, path)
            return os.path.normpath(path)

        # xhtml2pdf 0.2.21+ 的资源安全策略默认只允许读取工作目录，
        # 这里显式放行 md 所在目录与网络图片的临时目录（含局域网地址）
        try:
            from xhtml2pdf.config.resources import ResourceAccessPolicy, use_policy
            policy = ResourceAccessPolicy(
                base_dir=md_dir, extra_roots=(tmpdir,),
                allow_remote=True, allow_private_networks=True)
            ctx = use_policy(policy)
        except Exception:
            ctx = None

        try:
            if ctx:
                ctx.__enter__()
            with open(pdf_path, "wb") as f:
                status = pisa.CreatePDF(html, dest=f, path=md_dir + os.sep,
                                        encoding="utf-8", link_callback=link_callback)
        finally:
            if ctx:
                ctx.__exit__(None, None, None)
        if status.err:
            raise RuntimeError("PDF 生成失败，请检查 Markdown 内容（如特殊语法）")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def pdf_file_to_images(pdf_path, out_dir, dpi=400, fmt="png"):
    """单个 pdf → 多张图片，返回生成的图片路径列表"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open(pdf_path)
    stem = os.path.splitext(os.path.basename(pdf_path))[0]
    made = []
    try:
        for i, page in enumerate(doc, start=1):
            pix = page.get_pixmap(dpi=int(dpi))
            out = os.path.join(out_dir, "%s_p%02d.%s" % (stem, i, fmt.lower()))
            pix.save(out)
            made.append(out)
    finally:
        doc.close()
    return made


def batch_convert(kind, files, out_dir=None, dpi=400, img_fmt="png", log=print):
    """批量转换。out_dir 为 None 时输出到源文件所在目录。
    kind: 'docx' | 'pdf' | 'img'
    返回 (成功数, 失败数, 输出路径列表)
    """
    ok = fail = 0
    outs = []
    for idx, f in enumerate(files, 1):
        f = os.path.abspath(f)
        base = os.path.splitext(os.path.basename(f))[0]
        src_dir = os.path.dirname(f)
        log("[%d/%d] 正在转换：%s" % (idx, len(files), os.path.basename(f)))
        try:
            if kind == "docx":
                dest_dir = out_dir or src_dir
                dest = os.path.join(dest_dir, base + ".docx")
                md_file_to_docx(f, dest, log=log)
                outs.append(dest)
            elif kind == "pdf":
                dest_dir = out_dir or src_dir
                dest = os.path.join(dest_dir, base + ".pdf")
                md_file_to_pdf(f, dest, log=log)
                outs.append(dest)
            else:  # img
                dest_dir = out_dir or src_dir
                sub = os.path.join(dest_dir, base + "_图片")
                os.makedirs(sub, exist_ok=True)
                made = pdf_file_to_images(f, sub, dpi=dpi, fmt=img_fmt)
                outs.extend(made)
                dest = sub
            ok += 1
            log("    完成 -> %s" % dest)
        except Exception as e:
            fail += 1
            log("    失败：%s" % e)
            try:
                with open(ERROR_LOG, "a", encoding="utf-8") as ef:
                    ef.write("[%s] %s\n%s\n" % (os.path.basename(f), e, traceback.format_exc()))
            except Exception:
                pass
    return ok, fail, outs


# ---------------- 图形界面 ----------------

MD_PATTERNS = [("Markdown 文档", "*.md *.markdown"), ("所有文件", "*.*")]
PDF_PATTERNS = [("PDF 文档", "*.pdf"), ("所有文件", "*.*")]

TAB_CFG = {
    "md": dict(title="Markdown 转换",
               hint="把 .md 文件转换为 Word 或 PDF（可两种都要），图片自动嵌入。",
               patterns=MD_PATTERNS, exts=(".md", ".markdown")),
    "img": dict(title="PDF 转图片",
                hint="把 PDF 每一页导出为图片，可调清晰度。",
                patterns=PDF_PATTERNS, exts=(".pdf",)),
}

# MD 页的格式选择 → 实际转换目标列表
TARGET_CHOICE = {"word": ["docx"], "pdf": ["pdf"], "both": ["docx", "pdf"]}


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.configure(bg=C_WINDOW)
        self._apply_geometry()
        self._set_app_icon()

        self.log_q = queue.Queue()
        self.dnd_q = queue.Queue()   # 拖放回调 -> 界面的中转队列
        self.tabs = {}           # kind -> 状态字典
        self.panels = {}         # kind -> 面板 Frame
        self.nav_items = {}      # kind -> 导航项控件
        self.conv_btns = []      # (Button, 原文字)
        self._outline_btns = {}  # wrap Frame -> 内部 Button
        self.current = None

        self._init_theme()
        self._build_sidebar()
        self._build_main()
        self._build_statusbar()
        self._select("md")
        self._setup_dragdrop()
        self.after(100, self._poll_log)
        self.after(150, self._poll_dnd)

    # ---------- 基础 ----------

    def _apply_geometry(self):
        """窗口按 2:3（宽:高）比例，随系统缩放适配，完整放在任务栏以内的工作区"""
        self.update_idletasks()
        try:
            f = float(self.tk.call("tk", "scaling")) / 1.3333
        except Exception:
            f = 1.0
        self.ui_scale = max(1.0, f)
        sw = self.winfo_screenwidth()
        wa_w, wa_h = self._workarea()
        h = int(min(930 * self.ui_scale, wa_h * 0.94) * UI_SHRINK)
        w = h * 2 // 3
        x = max(0, (sw - w) // 2)
        y = max(0, (wa_h - h) // 3)
        self.geometry("%dx%d+%d+%d" % (w, h, x, y))
        self.minsize(int(w * 0.8), int(h * 0.8))
        # 映射后再应用一次，防止窗口管理器改位置
        self.after(80, lambda: self.geometry("%dx%d+%d+%d" % (w, h, x, y)))

    def _workarea(self):
        """Windows 工作区（去掉任务栏）；失败则退回整屏"""
        try:
            import ctypes

            class RECT(ctypes.Structure):
                _fields_ = [("left", ctypes.c_long), ("top", ctypes.c_long),
                            ("right", ctypes.c_long), ("bottom", ctypes.c_long)]
            rc = RECT()
            # SPI_GETWORKAREA
            if ctypes.windll.user32.SystemParametersInfoW(48, 0, ctypes.byref(rc), 0):
                return rc.right - rc.left, rc.bottom - rc.top
        except Exception:
            pass
        return self.winfo_screenwidth(), self.winfo_screenheight()

    def px(self, v):
        """按系统缩放换算像素（含整体缩小系数）"""
        return int(round(v * getattr(self, "ui_scale", 1.0) * UI_SHRINK))

    def fs(self, n):
        """按整体缩小系数换算字号（点）"""
        return max(7, int(round(n * UI_SHRINK)))

    def _set_app_icon(self):
        """用 Pillow 画一个程序图标（没有 Pillow 就跳过）"""
        try:
            from PIL import Image, ImageDraw, ImageFont, ImageTk
            img = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
            d = ImageDraw.Draw(img)
            d.rounded_rectangle([3, 3, 61, 61], radius=14, fill=(15, 108, 189, 255))
            try:
                font = ImageFont.truetype(os.path.join(FONT_DIR, "msyh.ttc"), 28)
            except Exception:
                font = ImageFont.truetype(os.path.join(FONT_DIR, "simhei.ttf"), 28)
            d.text((32, 33), "转", font=font, fill="white", anchor="mm")
            self._icon_img = ImageTk.PhotoImage(img)
            self.iconphoto(True, self._icon_img)
        except Exception:
            pass

    def _init_theme(self):
        style = ttk.Style(self)
        try:
            import sv_ttk
            sv_ttk.set_theme("light")
        except Exception:
            style.theme_use("clam")
        # 卡片内的 ttk 控件统一白底
        style.configure("Card.TRadiobutton", background=C_CARD, foreground=C_TEXT,
                        font=(FONT, self.fs(10)))
        style.map("Card.TRadiobutton", background=[("active", C_CARD)])

    # ---------- 小部件工厂 ----------

    def _card(self, parent, title):
        """白色卡片（1px 边框），返回 (外框, 内容区)。
        标题与内容分开容器，避免 pack/grid 混用冲突。"""
        outer = tk.Frame(parent, bg=C_BORDER)
        inner = tk.Frame(outer, bg=C_CARD)
        inner.pack(fill="both", expand=True, padx=1, pady=1)
        if title:
            tk.Label(inner, text=title, bg=C_CARD, fg=C_TEXT,
                     font=(FONT, self.fs(10), "bold"), anchor="w").pack(
                fill="x", padx=self.px(12), pady=(self.px(6), 2))
            body = tk.Frame(inner, bg=C_CARD)
            body.pack(fill="both", expand=True, padx=self.px(10), pady=(2, self.px(7)))
        else:
            body = tk.Frame(inner, bg=C_CARD)
            body.pack(fill="both", expand=True, padx=self.px(10), pady=self.px(6))
        return outer, body

    def _outline_btn(self, parent, text, command):
        """白底描边按钮"""
        wrap = tk.Frame(parent, bg=C_BORDER)
        b = tk.Button(wrap, text=text, command=command, bd=0, relief="flat",
                      bg=C_CARD, fg=C_TEXT, font=(FONT, self.fs(9)), cursor="hand2",
                      activebackground="#f0f0f0", activeforeground=C_TEXT,
                      disabledforeground="#a5a5a5",
                      padx=self.px(12), pady=self.px(4))
        b.pack(fill="both", expand=True, padx=1, pady=1)

        def hover(entered):
            if b["state"] == "normal":
                b.configure(bg="#f5f5f5" if entered else C_CARD)

        b.bind("<Enter>", lambda e: hover(True))
        b.bind("<Leave>", lambda e: hover(False))
        self._outline_btns[wrap] = b
        return wrap

    def _accent_btn(self, parent, text, command):
        """蓝色主按钮"""
        b = tk.Button(parent, text=text, command=command, bd=0, relief="flat",
                      bg=C_ACCENT, fg="white", font=(FONT, self.fs(11), "bold"),
                      activebackground=C_ACCENT_HOVER, activeforeground="white",
                      disabledforeground="#e2e2e2", cursor="hand2",
                      pady=self.px(8))

        def hover(entered):
            if b["state"] == "normal":
                b.configure(bg=C_ACCENT_HOVER if entered else C_ACCENT)

        b.bind("<Enter>", lambda e: hover(True))
        b.bind("<Leave>", lambda e: hover(False))
        return b

    def _soft_btn(self, parent, text, command):
        """浅蓝次要按钮"""
        b = tk.Button(parent, text=text, command=command, bd=0, relief="flat",
                      bg=C_ACCENT_LIGHT, fg=C_ACCENT, font=(FONT, self.fs(10), "bold"),
                      activebackground=C_ACCENT_SOFT_HOVER, activeforeground=C_ACCENT,
                      disabledforeground="#c3d4e5", cursor="hand2",
                      pady=self.px(6))

        def hover(entered):
            if b["state"] == "normal":
                b.configure(bg=C_ACCENT_SOFT_HOVER if entered else C_ACCENT_LIGHT)

        b.bind("<Enter>", lambda e: hover(True))
        b.bind("<Leave>", lambda e: hover(False))
        return b

    def _nav_item(self, parent, badge_text, label_text, command, row, kind=None):
        """侧栏导航项：小徽标 + 文字"""
        item = tk.Frame(parent, bg=C_SIDEBAR, cursor="hand2")
        item.grid(row=row, column=0, sticky="ew", padx=self.px(8), pady=1)
        badge = tk.Label(item, text=badge_text, width=2, bg="#eaeaea", fg="#5a5a5a",
                         font=(FONT, self.fs(9), "bold"))
        badge.pack(side="left", padx=(self.px(6), self.px(10)), pady=self.px(7))
        lbl = tk.Label(item, text=label_text, bg=C_SIDEBAR, fg="#3d3d3d",
                       font=(FONT, self.fs(10)))
        lbl.pack(side="left")

        def hover(entered):
            if kind is not None and kind == self.current:
                return
            bg = C_HOVER if entered else C_SIDEBAR
            item.configure(bg=bg)
            lbl.configure(bg=bg)

        for w in (item, badge, lbl):
            w.bind("<Button-1>", lambda e: command())
            w.bind("<Enter>", lambda e: hover(True))
            w.bind("<Leave>", lambda e: hover(False))
        return {"item": item, "badge": badge, "lbl": lbl}

    def _open_help(self):
        try:
            os.startfile(os.path.join(APP_DIR, "使用说明.txt"))
        except Exception:
            messagebox.showinfo(APP_TITLE, "找不到使用说明.txt")

    # ---------- 布局 ----------

    def _build_sidebar(self):
        sb = tk.Frame(self, bg=C_SIDEBAR, width=self.px(140))
        sb.grid(row=0, column=0, sticky="nsw")
        sb.grid_propagate(False)
        sb.columnconfigure(0, weight=1)
        sb.rowconfigure(4, weight=1)
        # 右侧 1px 分隔线（用 grid 放在独立列，避免 pack/grid 混用）
        tk.Frame(sb, bg=C_BORDER, width=1).grid(row=0, column=1, rowspan=7, sticky="ns")

        brand = tk.Frame(sb, bg=C_SIDEBAR)
        brand.grid(row=0, column=0, sticky="ew", padx=self.px(6), pady=(self.px(14), self.px(10)))
        tk.Label(brand, text="文档转换工具", bg=C_SIDEBAR, fg=C_TEXT,
                 font=(FONT, self.fs(11), "bold"), anchor="w").pack(fill="x")
        tk.Label(brand, text="MD·Word·PDF·图片", bg=C_SIDEBAR, fg=C_SUB,
                 font=(FONT, self.fs(8)), anchor="w").pack(fill="x", pady=(2, 0))

        navs = (("md", "MD", "MD 转换"),
                ("img", "图", "PDF 转图片"))
        for i, (kind, badge, text) in enumerate(navs, start=1):
            self.nav_items[kind] = self._nav_item(
                sb, badge, text, lambda k=kind: self._select(k), row=i, kind=kind)

        self._nav_item(sb, "?", "使用说明", self._open_help, row=5)

        tk.Label(sb, text="v1.9 · 免费离线", bg=C_SIDEBAR, fg="#a5a5a5",
                 font=(FONT, self.fs(8))).grid(row=6, column=0, sticky="w",
                                      padx=self.px(16), pady=(6, self.px(12)))

    def _build_main(self):
        main = tk.Frame(self, bg=C_WINDOW)
        main.grid(row=0, column=1, sticky="nsew")
        main.columnconfigure(0, weight=1)
        main.rowconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)
        self.columnconfigure(1, weight=1)

        container = tk.Frame(main, bg=C_WINDOW)
        container.grid(row=0, column=0, sticky="nsew",
                       padx=self.px(18), pady=(self.px(12), self.px(6)))
        container.columnconfigure(0, weight=1)
        container.rowconfigure(0, weight=1)

        for kind in ("md", "img"):
            pf = tk.Frame(container, bg=C_WINDOW)
            pf.grid(row=0, column=0, sticky="nsew")
            pf.columnconfigure(0, weight=1)
            self.panels[kind] = pf
            self._build_panel(kind, pf)

        self._build_log(main)

    def _build_panel(self, kind, parent):
        cfg = TAB_CFG[kind]
        st = self.tabs.setdefault(kind, {})

        tk.Label(parent, text=cfg["title"], bg=C_WINDOW, fg=C_TEXT,
                 font=(FONT, self.fs(15), "bold"), anchor="w").grid(row=0, column=0, sticky="w")
        tk.Label(parent, text=cfg["hint"], bg=C_WINDOW, fg=C_SUB,
                 font=(FONT, self.fs(9)), anchor="w", justify="left",
                 wraplength=self.px(295)).grid(row=1, column=0, sticky="w",
                                               pady=(2, self.px(7)))

        # ① 文件卡片
        card, inner = self._card(parent, "① 选择文件（可多选）")
        card.grid(row=2, column=0, sticky="ew")

        lbf = tk.Frame(inner, bg=C_CARD)
        lbf.pack(fill="both", expand=True)
        lb = tk.Listbox(lbf, bd=0, relief="flat", highlightthickness=1,
                        highlightbackground=C_BORDER, highlightcolor=C_ACCENT,
                        bg=C_CARD, fg=C_TEXT, selectbackground=C_ACCENT_LIGHT,
                        selectforeground=C_TEXT, activestyle="none",
                        font=(FONT, self.fs(9)), height=5)
        sb = ttk.Scrollbar(lbf, orient="vertical", command=lb.yview)
        lb.configure(yscrollcommand=sb.set)
        lb.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        st["list"] = lb
        st["paths"] = []
        st["hint"] = tk.Label(lb, text="还没有文件，可点击「＋ 添加文件…」或直接拖入文件",
                              bg=C_CARD, fg="#9a9a9a", font=(FONT, self.fs(9)))

        row = tk.Frame(inner, bg=C_CARD)
        row.pack(fill="x", pady=(self.px(6), 0))
        self._outline_btn(row, "＋ 添加文件…",
                          lambda k=kind: self._add_files(k)).pack(side="left")
        self._outline_btn(row, "移除选中",
                          lambda k=kind: self._remove_sel(k)).pack(side="left", padx=(8, 0))
        self._outline_btn(row, "清空",
                          lambda k=kind: self._clear(k)).pack(side="left", padx=(8, 0))

        # ② 输出位置卡片（路径框独占一行，保证足够宽）
        card2, inner2 = self._card(parent, "② 输出位置")
        card2.grid(row=3, column=0, sticky="ew", pady=(self.px(12), 0))
        inner2.columnconfigure(0, weight=1)
        st["same_var"] = tk.BooleanVar(value=True)
        ttk.Radiobutton(inner2, text="保存到源文件所在文件夹", value=True,
                        variable=st["same_var"], style="Card.TRadiobutton",
                        command=lambda k=kind: self._out_mode(k)
                        ).grid(row=0, column=0, columnspan=2, sticky="w", pady=1)
        ttk.Radiobutton(inner2, text="保存到指定文件夹：", value=False,
                        variable=st["same_var"], style="Card.TRadiobutton",
                        command=lambda k=kind: self._out_mode(k)
                        ).grid(row=1, column=0, columnspan=2, sticky="w", pady=1)
        st["out_var"] = tk.StringVar(
            value=os.path.join(os.path.expanduser("~"), "Desktop"))
        st["out_entry"] = ttk.Entry(inner2, textvariable=st["out_var"],
                                    font=(FONT, self.fs(12)))
        st["out_entry"].grid(row=2, column=0, sticky="ew", pady=1)
        st["out_btn"] = self._outline_btn(inner2, "浏览…",
                                          lambda k=kind: self._pick_outdir(k))
        st["out_btn"].grid(row=2, column=1, padx=(6, 0), pady=1)
        self._out_mode(kind)

        # ③ 转换格式（MD 页）/ 图片选项（PDF 转图片页）
        if kind == "md":
            card3, inner3 = self._card(parent, "③ 转换格式")
            card3.grid(row=4, column=0, sticky="ew", pady=(self.px(12), 0))
            st["target_var"] = tk.StringVar(value="both")
            for val, text in (("word", ".docx"),
                              ("pdf", ".pdf"),
                              ("both", "Both")):
                ttk.Radiobutton(inner3, text=text, value=val,
                                variable=st["target_var"],
                                style="Card.TRadiobutton").pack(
                    side="left", padx=(0, self.px(18)), pady=(0, 2))
            btn_row = 5
        else:
            card3, inner3 = self._card(parent, "③ 图片选项")
            card3.grid(row=4, column=0, sticky="ew", pady=(self.px(12), 0))
            f = tk.Frame(inner3, bg=C_CARD)
            f.pack(fill="x")
            tk.Label(f, text="清晰度 (DPI)", bg=C_CARD, fg=C_SUB,
                     font=(FONT, self.fs(9))).pack(side="left")
            st["dpi_var"] = tk.StringVar(value="400")
            ttk.Combobox(f, textvariable=st["dpi_var"],
                         values=("96", "150", "200", "300", "400"),
                         state="readonly", width=5, font=(FONT, self.fs(10))
                         ).pack(side="left", padx=(6, 14))
            tk.Label(f, text="格式", bg=C_CARD, fg=C_SUB,
                     font=(FONT, self.fs(9))).pack(side="left")
            st["fmt_var"] = tk.StringVar(value="png")
            ttk.Combobox(f, textvariable=st["fmt_var"], values=("png", "jpg"),
                         state="readonly", width=5, font=(FONT, self.fs(10))
                         ).pack(side="left", padx=(6, 0))
            btn_row = 5

        # 转换按钮
        text = "开始转换" if kind == "md" else "转换为图片"
        b = self._accent_btn(parent, text, lambda k=kind: self._convert(k))
        b.grid(row=btn_row, column=0, sticky="ew", pady=(self.px(10), 0))
        self.conv_btns.append((b, text))
        self._refresh_list(kind)

    def _build_log(self, main):
        card, inner = self._card(main, None)
        card.grid(row=1, column=0, sticky="ew", pady=(self.px(10), 0))

        hdr = tk.Frame(inner, bg=C_CARD)
        hdr.pack(fill="x", pady=(2, 0))
        tk.Label(hdr, text="转换日志", bg=C_CARD, fg=C_SUB,
                 font=(FONT, self.fs(9), "bold")).pack(side="left")
        clear = tk.Button(hdr, text="清空", bd=0, bg=C_CARD, fg=C_SUB,
                          font=(FONT, self.fs(8)), cursor="hand2",
                          activebackground=C_CARD, activeforeground=C_TEXT,
                          command=self._clear_log)
        clear.pack(side="right")

        body = tk.Frame(inner, bg=C_CARD)
        body.pack(fill="x", pady=(2, 2))
        self.log_text = tk.Text(body, bd=0, relief="flat", bg=C_CARD, fg=C_TEXT,
                                font=(FONT, self.fs(9)), height=5, wrap="word",
                                state="disabled", padx=4, pady=2,
                                insertbackground=C_TEXT)
        lsb = ttk.Scrollbar(body, orient="vertical", command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=lsb.set)
        lsb.pack(side="right", fill="y")
        self.log_text.pack(side="left", fill="both", expand=True)
        for tag, color in (("ok", C_OK), ("err", C_ERR), ("info", "#4a4a4a")):
            self.log_text.tag_configure(tag, foreground=color)
        self.log_text.tag_configure("head", foreground=C_ACCENT,
                                    font=(FONT, self.fs(9), "bold"))

    def _build_statusbar(self):
        status = tk.Frame(self, bg=C_WINDOW)
        status.grid(row=1, column=0, columnspan=2, sticky="ew")
        tk.Frame(status, bg=C_BORDER, height=1).pack(fill="x")

        row = tk.Frame(status, bg=C_WINDOW)
        row.pack(fill="x")
        self.dot = tk.Label(row, text="●", fg="#8a8886", bg=C_WINDOW, font=(FONT, self.fs(9)))
        self.dot.pack(side="left", padx=(self.px(12), 4), pady=4)
        self.status_var = tk.StringVar(value="就绪")
        tk.Label(row, textvariable=self.status_var, bg=C_WINDOW, fg=C_SUB,
                 font=(FONT, self.fs(9))).pack(side="left")
        self.progress = ttk.Progressbar(row, mode="indeterminate", length=110)

        tk.Label(row, text="免费离线 · 所有转换都在本机完成，文件不经过网络",
                 bg=C_WINDOW, fg="#a5a5a5", font=(FONT, self.fs(8))).pack(
            side="right", padx=self.px(14), pady=4)

    # ---------- 导航切换 ----------

    def _select(self, kind):
        self.current = kind
        for k, d in self.nav_items.items():
            if k == kind:
                d["item"].configure(bg=C_ACCENT_LIGHT)
                d["lbl"].configure(bg=C_ACCENT_LIGHT, fg=C_ACCENT)
                d["badge"].configure(bg=C_ACCENT, fg="white")
            else:
                d["item"].configure(bg=C_SIDEBAR)
                d["lbl"].configure(bg=C_SIDEBAR, fg="#3d3d3d")
                d["badge"].configure(bg="#eaeaea", fg="#5a5a5a")
        for k, p in self.panels.items():
            if k == kind:
                p.grid()
            else:
                p.grid_remove()

    # ---------- 拖放支持 ----------

    def _setup_dragdrop(self):
        """把整个窗口注册为文件拖放目标：文件可拖到窗口任意位置，
        会加入当前页面的列表；拖入文件夹时取其中匹配的文件。
        使用 Windows 原生 WM_DROPFILES 机制（windnd）。
        注意：系统回调里不能直接操作 Tk 控件（GIL/Tcl 状态限制），
        所以这里只把路径放进队列，由 _poll_dnd 在正常状态更新界面。"""
        if not _HAS_DND:
            return

        def on_drop(files):
            try:
                paths = []
                for b in files:
                    if isinstance(b, bytes):
                        paths.append(b.decode("gbk", errors="replace"))
                    else:
                        paths.append(str(b))
                if paths:
                    self.dnd_q.put(paths)
            except Exception:
                try:
                    with open(ERROR_LOG, "a", encoding="utf-8") as f:
                        f.write("[拖放回调] \n%s\n" % traceback.format_exc())
                except Exception:
                    pass

        try:
            windnd.hook_dropfiles(self, func=on_drop, force_unicode=True)
        except Exception:
            self._log("拖放功能初始化失败（不影响其他功能），详见 error.log")
            try:
                with open(ERROR_LOG, "a", encoding="utf-8") as f:
                    f.write("[拖放初始化] \n%s\n" % traceback.format_exc())
            except Exception:
                pass

    def _poll_dnd(self):
        """处理拖放队列：真正加入列表、刷新界面（在主循环安全时机执行）"""
        try:
            while True:
                paths = self.dnd_q.get_nowait()
                added, skipped = self._add_dropped(self.current, paths)
                self._refresh_list(self.current)
                if added or skipped:
                    msg = "拖入：%d 个文件已添加" % added
                    if skipped:
                        msg += "，跳过 %d 个不支持的" % skipped
                    self._log(msg)
                else:
                    self._log("拖入的内容中没有可识别的文件")
        except queue.Empty:
            pass
        self.after(150, self._poll_dnd)

    def report_callback_exception(self, exc, val, tb):
        """界面回调里的异常写入 error.log（pythonw 无控制台，否则看不见）"""
        try:
            with open(ERROR_LOG, "a", encoding="utf-8") as f:
                f.write("[界面回调] %r\n%s\n" % (val, "".join(traceback.format_tb(tb))))
        except Exception:
            pass

    def _add_dropped(self, kind, paths):
        """把路径加入指定页面的列表，自动去重、按扩展名过滤。
        拖入文件夹时，取文件夹下一层中匹配的文件。
        返回 (新增数, 跳过数)"""
        st = self.tabs.get(kind)
        if st is None:
            return 0, 0
        exts = TAB_CFG[kind]["exts"]
        added = skipped = 0
        for p in paths:
            p = p.strip().strip("{}")
            if os.path.isdir(p):
                try:
                    names = sorted(os.listdir(p))
                except OSError:
                    continue
                for name in names:
                    fp = os.path.join(p, name)
                    if os.path.isfile(fp) and name.lower().endswith(exts):
                        if fp not in st["paths"]:
                            st["paths"].append(fp)
                            added += 1
            elif os.path.isfile(p):
                if p.lower().endswith(exts):
                    if p not in st["paths"]:
                        st["paths"].append(p)
                        added += 1
                else:
                    skipped += 1
        return added, skipped

    # ---------- 文件列表 ----------

    def _add_files(self, kind):
        cfg = TAB_CFG[kind]
        names = filedialog.askopenfilenames(title="选择文件", filetypes=cfg["patterns"])
        self._add_dropped(kind, list(names))
        self._refresh_list(kind)

    def _remove_sel(self, kind):
        st = self.tabs[kind]
        for i in reversed(st["list"].curselection()):
            del st["paths"][i]
        self._refresh_list(kind)

    def _clear(self, kind):
        self.tabs[kind]["paths"] = []
        self._refresh_list(kind)

    def _refresh_list(self, kind):
        st = self.tabs[kind]
        lb = st["list"]
        lb.delete(0, "end")
        for p in st["paths"]:
            lb.insert("end", p)
        if st["paths"]:
            st["hint"].place_forget()
        else:
            st["hint"].place(relx=0.5, rely=0.38, anchor="center")

    def _pick_outdir(self, kind):
        d = filedialog.askdirectory(title="选择输出文件夹")
        if d:
            self.tabs[kind]["out_var"].set(d.replace("/", "\\"))

    def _out_mode(self, kind):
        st = self.tabs[kind]
        spec = not st["same_var"].get()
        st["out_entry"].configure(state="normal" if spec else "disabled")
        b = self._outline_btns[st["out_btn"]]
        b.configure(state="normal" if spec else "disabled")

    # ---------- 转换 ----------

    def _convert(self, state_kind):
        """state_kind：取哪个面板的文件列表/输出设置；
        转换目标由面板内选择决定（MD 页的 1/2/3 单选）。"""
        st = self.tabs[state_kind]
        files = list(st["paths"])
        if not files:
            messagebox.showwarning(APP_TITLE, "请先点击「＋ 添加文件…」选择要转换的文件。")
            return
        out_dir = None
        if not st["same_var"].get():
            out_dir = st["out_var"].get().strip()
            if not out_dir:
                messagebox.showwarning(APP_TITLE, "请先选择输出文件夹。")
                return
            os.makedirs(out_dir, exist_ok=True)
        if state_kind == "img":
            targets = ["img"]
            dpi = int(st["dpi_var"].get())
            img_fmt = st["fmt_var"].get()
        else:
            choice = st["target_var"].get()
            targets = list(TARGET_CHOICE.get(choice, TARGET_CHOICE["both"]))
            dpi, img_fmt = 150, "png"

        self._set_busy(True)
        name = " + ".join(TARGET_NAME[t] for t in targets)
        self._log("—— 开始转换（%s），共 %d 个文件 ——" % (name, len(files)))
        threading.Thread(target=self._worker,
                         args=(targets, files, out_dir, dpi, img_fmt),
                         daemon=True).start()

    def _worker(self, targets, files, out_dir, dpi, fmt):
        ok = fail = 0
        outs = []
        try:
            for t in targets:
                if len(targets) > 1:
                    self._log("—— 转换为 %s ——" % TARGET_NAME[t])
                o, f, out = batch_convert(t, files, out_dir, dpi, fmt, log=self._log)
                ok += o
                fail += f
                outs.extend(out)
        except Exception:
            self._log("发生意外错误：\n" + traceback.format_exc())
        self.after(0, self._done, ok, fail, outs)

    def _done(self, ok, fail, outs):
        self._set_busy(False)
        msg = "完成：成功 %d 个，失败 %d 个" % (ok, fail)
        self.status_var.set(msg)
        self.dot.configure(fg="#13a10e" if fail == 0 else
                           ("#c50f1f" if ok == 0 else "#ffb900"))
        self._log("—— %s ——" % msg)
        if ok and outs:
            folder = outs[0] if os.path.isdir(outs[0]) else os.path.dirname(outs[0])
            if messagebox.askyesno(APP_TITLE, "%s\n\n是否打开输出文件夹？" % msg):
                try:
                    os.startfile(folder)
                except Exception:
                    pass
        elif fail:
            messagebox.showerror(
                APP_TITLE, "全部转换失败（%d 个），详情见日志窗口。\n"
                           "若提示缺少库，请先双击工具文件夹里的「安装依赖.bat」。" % fail)

    def _set_busy(self, busy):
        for b, text in self.conv_btns:
            if busy:
                b.configure(state="disabled", text="转换中…")
            else:
                b.configure(state="normal", text=text)
        if busy:
            self.dot.configure(fg="#ffb900")
            self.status_var.set("转换中，请稍候…")
            self.progress.pack(side="left", padx=(8, 0))
            self.progress.start(14)
        else:
            self.progress.stop()
            self.progress.pack_forget()

    # ---------- 日志（线程安全） ----------

    def _log(self, text):
        self.log_q.put(text)

    def _poll_log(self):
        try:
            while True:
                line = self.log_q.get_nowait()
                if line.startswith("——"):
                    tag = "head"
                elif "完成 ->" in line:
                    tag = "ok"
                elif "失败" in line or "错误" in line:
                    tag = "err"
                else:
                    tag = "info"
                self.log_text.configure(state="normal")
                self.log_text.insert("end", line + "\n", tag)
                self.log_text.see("end")
                self.log_text.configure(state="disabled")
        except queue.Empty:
            pass
        self.after(100, self._poll_log)

    def _clear_log(self):
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")


def main():
    try:
        # 让界面在高分屏上清晰
        import ctypes
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    except Exception:
        pass
    try:
        App().mainloop()
    except Exception:
        err = traceback.format_exc()
        try:
            with open(ERROR_LOG, "a", encoding="utf-8") as f:
                f.write("[启动] \n%s\n" % err)
        except Exception:
            pass
        try:
            root = tk.Tk()
            root.withdraw()
            messagebox.showerror(APP_TITLE, "启动失败，详情见 %s\n\n%s" % (ERROR_LOG, err))
        except Exception:
            pass


if __name__ == "__main__":
    main()
